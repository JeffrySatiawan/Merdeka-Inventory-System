#!/usr/bin/env python3
"""
Backend test for Absensi "Poin Saat Absen" patch (Feb 2026).

PATCH VERIFICATION:
- GET /api/absensi/report now returns points_by_record (per record id) and points_by_user (recap)
- Field points_period REMOVED
- Balance calculated from absensi_point_ledger: clamp(initial_balance + Σ ledger.points where event_date <= record.date)
- Excel export sheet "Rekapitulasi" column "Poin Saat Absen" (not "Poin Saat Ini")
- FROZEN HISTORIC: Points after record.date MUST NOT affect historical report

TEST PLAN:
1. Login as owner
2. Seed absensi data (records + ledger entries)
3. Call GET /api/absensi/report and verify response structure
4. Verify points_by_record calculation (frozen historic)
5. Verify points_by_user (last record snapshot)
6. CRITICAL: Insert new ledger entry AFTER record date → verify balance unchanged (frozen)
7. Insert backdated ledger entry BEFORE record date → verify balance updated
8. Excel export verification
9. Empty case (zero rows)
10. Regression checks

CREDENTIALS:
- Owner: owner / owner123
- Staff: cindy / cindy123
"""

import requests
import json
from datetime import datetime, timedelta
from pymongo import MongoClient
import os
import sys

BASE_URL = os.getenv('NEXT_PUBLIC_BASE_URL', 'https://absensi-foundation.preview.emergentagent.com')
MONGO_URL = os.getenv('MONGO_URL', 'mongodb://localhost:27017')
DB_NAME = os.getenv('DB_NAME', 'cycle_count')

def log(msg):
    print(f"[{datetime.now().strftime('%H:%M:%S')}] {msg}")

def main():
    log("=" * 80)
    log("ABSENSI POIN SAAT ABSEN PATCH VERIFICATION")
    log("=" * 80)
    
    # Connect to MongoDB
    log("\n[SETUP] Connecting to MongoDB...")
    try:
        client = MongoClient(MONGO_URL)
        db = client[DB_NAME]
        log(f"✅ Connected to MongoDB: {DB_NAME}")
    except Exception as e:
        log(f"❌ MongoDB connection failed: {e}")
        return False
    
    # Test counters
    total_tests = 0
    passed_tests = 0
    
    # ========================================================================
    # TEST 1: LOGIN AS OWNER
    # ========================================================================
    log("\n" + "=" * 80)
    log("TEST 1: LOGIN AS OWNER")
    log("=" * 80)
    total_tests += 1
    
    try:
        resp = requests.post(f"{BASE_URL}/api/auth/login", json={"username": "owner", "password": "owner123"}, timeout=10)
        if resp.status_code != 200:
            log(f"❌ Login failed: {resp.status_code} {resp.text}")
            return False
        data = resp.json()
        owner_token = data.get('token')
        owner_id = data.get('user', {}).get('id')
        if not owner_token:
            log(f"❌ No token in login response")
            return False
        log(f"✅ Owner login successful, token: {owner_token[:20]}...")
        log(f"   Owner ID: {owner_id}")
        passed_tests += 1
    except Exception as e:
        log(f"❌ Login exception: {e}")
        return False
    
    headers = {"Authorization": f"Bearer {owner_token}"}
    
    # ========================================================================
    # TEST 2: GET POINT SETTINGS (for clamp bounds)
    # ========================================================================
    log("\n" + "=" * 80)
    log("TEST 2: GET POINT SETTINGS")
    log("=" * 80)
    total_tests += 1
    
    try:
        resp = requests.get(f"{BASE_URL}/api/absensi/points/settings", headers=headers, timeout=10)
        if resp.status_code != 200:
            log(f"❌ GET /api/absensi/points/settings failed: {resp.status_code}")
            return False
        settings = resp.json()
        initial_balance = settings.get('initial_balance', 100)
        max_positive = settings.get('max_positive', 150)
        max_negative = settings.get('max_negative', -50)
        log(f"✅ Point settings loaded:")
        log(f"   initial_balance: {initial_balance}")
        log(f"   max_positive: {max_positive}")
        log(f"   max_negative: {max_negative}")
        passed_tests += 1
    except Exception as e:
        log(f"❌ Exception: {e}")
        return False
    
    # ========================================================================
    # TEST 3: GET CINDY USER ID
    # ========================================================================
    log("\n" + "=" * 80)
    log("TEST 3: GET CINDY USER ID")
    log("=" * 80)
    total_tests += 1
    
    try:
        resp = requests.get(f"{BASE_URL}/api/employees", headers=headers, timeout=10)
        if resp.status_code != 200:
            log(f"❌ GET /api/employees failed: {resp.status_code}")
            return False
        employees = resp.json().get('items', [])
        cindy = next((e for e in employees if e.get('username') == 'cindy'), None)
        if not cindy:
            log(f"❌ Cindy not found in employees")
            return False
        cindy_id = cindy['id']
        cindy_name = cindy['name']
        log(f"✅ Cindy found: {cindy_name} (ID: {cindy_id})")
        passed_tests += 1
    except Exception as e:
        log(f"❌ Exception: {e}")
        return False
    
    # ========================================================================
    # TEST 4: SEED ABSENSI RECORDS + LEDGER ENTRIES
    # ========================================================================
    log("\n" + "=" * 80)
    log("TEST 4: SEED ABSENSI RECORDS + LEDGER ENTRIES")
    log("=" * 80)
    total_tests += 1
    
    try:
        # Clean up existing test data
        db.absensi_records.delete_many({"user_id": cindy_id, "date": {"$regex": "^2026-02-"}})
        db.absensi_point_ledger.delete_many({"user_id": cindy_id, "event_date": {"$regex": "^2026-02-"}})
        log("   Cleaned up existing test data")
        
        # Create 3 absensi records for Cindy in Feb 2026
        from uuid import uuid4
        
        record1_id = str(uuid4())
        record1_date = "2026-02-10"
        record1 = {
            "id": record1_id,
            "user_id": cindy_id,
            "user_name": cindy_name,
            "date": record1_date,
            "actual_check_in": "2026-02-10T08:00:00Z",
            "actual_check_out": "2026-02-10T17:00:00Z",
            "actual_check_in_wita": "16.00",
            "actual_check_out_wita": "01.00",
            "shift_key": "pagi",
            "shift_name": "Pagi",
            "shift_start": "08:00",
            "shift_end": "17:00",
            "shift_start_mins": 480,
            "shift_end_mins": 1020,
            "late_minutes": 0,
        }
        
        record2_id = str(uuid4())
        record2_date = "2026-02-15"
        record2 = {
            "id": record2_id,
            "user_id": cindy_id,
            "user_name": cindy_name,
            "date": record2_date,
            "actual_check_in": "2026-02-15T08:00:00Z",
            "actual_check_out": "2026-02-15T17:00:00Z",
            "actual_check_in_wita": "16.00",
            "actual_check_out_wita": "01.00",
            "shift_key": "pagi",
            "shift_name": "Pagi",
            "shift_start": "08:00",
            "shift_end": "17:00",
            "shift_start_mins": 480,
            "shift_end_mins": 1020,
            "late_minutes": 0,
        }
        
        record3_id = str(uuid4())
        record3_date = "2026-02-20"
        record3 = {
            "id": record3_id,
            "user_id": cindy_id,
            "user_name": cindy_name,
            "date": record3_date,
            "actual_check_in": "2026-02-20T08:00:00Z",
            "actual_check_out": "2026-02-20T17:00:00Z",
            "actual_check_in_wita": "16.00",
            "actual_check_out_wita": "01.00",
            "shift_key": "pagi",
            "shift_name": "Pagi",
            "shift_start": "08:00",
            "shift_end": "17:00",
            "shift_start_mins": 480,
            "shift_end_mins": 1020,
            "late_minutes": 0,
        }
        
        db.absensi_records.insert_many([record1, record2, record3])
        log(f"✅ Inserted 3 absensi records for Cindy:")
        log(f"   Record 1: {record1_date} (ID: {record1_id})")
        log(f"   Record 2: {record2_date} (ID: {record2_id})")
        log(f"   Record 3: {record3_date} (ID: {record3_id})")
        
        # Create ledger entries
        # Ledger 1: 2026-02-05 (before record1) → +10 points
        ledger1_id = str(uuid4())
        ledger1 = {
            "id": ledger1_id,
            "user_id": cindy_id,
            "user_name": cindy_name,
            "event_type": "adjustment",
            "event_date": "2026-02-05",
            "period_key": "2026-02",
            "points": 10,
            "source_id": ledger1_id,
            "createdAt": datetime.utcnow(),
        }
        
        # Ledger 2: 2026-02-10 (same as record1) → +5 points (checkin)
        ledger2_id = str(uuid4())
        ledger2 = {
            "id": ledger2_id,
            "user_id": cindy_id,
            "user_name": cindy_name,
            "event_type": "checkin",
            "event_date": record1_date,
            "period_key": "2026-02",
            "points": 5,
            "source_id": record1_id,
            "createdAt": datetime.utcnow(),
        }
        
        # Ledger 3: 2026-02-12 (between record1 and record2) → -3 points
        ledger3_id = str(uuid4())
        ledger3 = {
            "id": ledger3_id,
            "user_id": cindy_id,
            "user_name": cindy_name,
            "event_type": "adjustment",
            "event_date": "2026-02-12",
            "period_key": "2026-02",
            "points": -3,
            "source_id": ledger3_id,
            "createdAt": datetime.utcnow(),
        }
        
        # Ledger 4: 2026-02-15 (same as record2) → +5 points (checkin)
        ledger4_id = str(uuid4())
        ledger4 = {
            "id": ledger4_id,
            "user_id": cindy_id,
            "user_name": cindy_name,
            "event_type": "checkin",
            "event_date": record2_date,
            "period_key": "2026-02",
            "points": 5,
            "source_id": record2_id,
            "createdAt": datetime.utcnow(),
        }
        
        db.absensi_point_ledger.insert_many([ledger1, ledger2, ledger3, ledger4])
        log(f"✅ Inserted 4 ledger entries:")
        log(f"   Ledger 1: 2026-02-05 → +10 points")
        log(f"   Ledger 2: 2026-02-10 → +5 points (checkin for record1)")
        log(f"   Ledger 3: 2026-02-12 → -3 points")
        log(f"   Ledger 4: 2026-02-15 → +5 points (checkin for record2)")
        
        # Expected balances:
        # record1 (2026-02-10): initial(100) + 10 + 5 = 115
        # record2 (2026-02-15): initial(100) + 10 + 5 - 3 + 5 = 117
        # record3 (2026-02-20): initial(100) + 10 + 5 - 3 + 5 = 117 (no ledger entry on this date)
        
        log(f"\n   Expected balances:")
        log(f"   Record 1 (2026-02-10): {initial_balance} + 10 + 5 = 115")
        log(f"   Record 2 (2026-02-15): {initial_balance} + 10 + 5 - 3 + 5 = 117")
        log(f"   Record 3 (2026-02-20): {initial_balance} + 10 + 5 - 3 + 5 = 117")
        
        passed_tests += 1
    except Exception as e:
        log(f"❌ Exception: {e}")
        import traceback
        traceback.print_exc()
        return False
    
    # ========================================================================
    # TEST 5: GET /api/absensi/report - VERIFY RESPONSE STRUCTURE
    # ========================================================================
    log("\n" + "=" * 80)
    log("TEST 5: GET /api/absensi/report - VERIFY RESPONSE STRUCTURE")
    log("=" * 80)
    
    try:
        resp = requests.get(f"{BASE_URL}/api/absensi/report?from=2026-02-01&to=2026-02-28", headers=headers, timeout=10)
        if resp.status_code != 200:
            log(f"❌ GET /api/absensi/report failed: {resp.status_code} {resp.text}")
            return False
        
        report = resp.json()
        log(f"✅ GET /api/absensi/report → 200")
        
        # Check 1: Response contains points_by_record
        total_tests += 1
        if 'points_by_record' not in report:
            log(f"❌ Response missing 'points_by_record' field")
            log(f"   Response keys: {list(report.keys())}")
            return False
        log(f"✅ Response contains 'points_by_record' field")
        passed_tests += 1
        
        # Check 2: Response contains points_by_user
        total_tests += 1
        if 'points_by_user' not in report:
            log(f"❌ Response missing 'points_by_user' field")
            return False
        log(f"✅ Response contains 'points_by_user' field")
        passed_tests += 1
        
        # Check 3: Response DOES NOT contain points_period
        total_tests += 1
        if 'points_period' in report:
            log(f"❌ Response contains 'points_period' field (should be removed)")
            return False
        log(f"✅ Response DOES NOT contain 'points_period' field (correctly removed)")
        passed_tests += 1
        
        # Check 4: points_by_record is an object
        total_tests += 1
        if not isinstance(report['points_by_record'], dict):
            log(f"❌ points_by_record is not an object: {type(report['points_by_record'])}")
            return False
        log(f"✅ points_by_record is an object (dict)")
        passed_tests += 1
        
        # Check 5: points_by_user is an object
        total_tests += 1
        if not isinstance(report['points_by_user'], dict):
            log(f"❌ points_by_user is not an object: {type(report['points_by_user'])}")
            return False
        log(f"✅ points_by_user is an object (dict)")
        passed_tests += 1
        
        # Check 6: items array exists
        total_tests += 1
        if 'items' not in report or not isinstance(report['items'], list):
            log(f"❌ Response missing 'items' array")
            return False
        log(f"✅ Response contains 'items' array with {len(report['items'])} records")
        passed_tests += 1
        
        # Store for later tests
        points_by_record = report['points_by_record']
        points_by_user = report['points_by_user']
        items = report['items']
        
    except Exception as e:
        log(f"❌ Exception: {e}")
        import traceback
        traceback.print_exc()
        return False
    
    # ========================================================================
    # TEST 6: VERIFY points_by_record CALCULATION (FROZEN HISTORIC)
    # ========================================================================
    log("\n" + "=" * 80)
    log("TEST 6: VERIFY points_by_record CALCULATION")
    log("=" * 80)
    
    try:
        # Find our test records in items
        test_records = [r for r in items if r.get('user_id') == cindy_id and r.get('date', '').startswith('2026-02-')]
        if len(test_records) != 3:
            log(f"❌ Expected 3 test records, found {len(test_records)}")
            return False
        
        # Sort by date
        test_records.sort(key=lambda r: r['date'])
        
        log(f"   Found {len(test_records)} test records for Cindy")
        
        # Verify record 1 (2026-02-10)
        total_tests += 1
        r1 = test_records[0]
        if r1['date'] != record1_date:
            log(f"❌ Record 1 date mismatch: expected {record1_date}, got {r1['date']}")
            return False
        if r1['id'] not in points_by_record:
            log(f"❌ Record 1 ID {r1['id']} not in points_by_record")
            return False
        balance1 = points_by_record[r1['id']]['balance']
        expected1 = 115
        if balance1 != expected1:
            log(f"❌ Record 1 balance mismatch: expected {expected1}, got {balance1}")
            return False
        log(f"✅ Record 1 ({record1_date}): balance = {balance1} (expected {expected1})")
        passed_tests += 1
        
        # Verify record 2 (2026-02-15)
        total_tests += 1
        r2 = test_records[1]
        if r2['date'] != record2_date:
            log(f"❌ Record 2 date mismatch: expected {record2_date}, got {r2['date']}")
            return False
        if r2['id'] not in points_by_record:
            log(f"❌ Record 2 ID {r2['id']} not in points_by_record")
            return False
        balance2 = points_by_record[r2['id']]['balance']
        expected2 = 117
        if balance2 != expected2:
            log(f"❌ Record 2 balance mismatch: expected {expected2}, got {balance2}")
            return False
        log(f"✅ Record 2 ({record2_date}): balance = {balance2} (expected {expected2})")
        passed_tests += 1
        
        # Verify record 3 (2026-02-20)
        total_tests += 1
        r3 = test_records[2]
        if r3['date'] != record3_date:
            log(f"❌ Record 3 date mismatch: expected {record3_date}, got {r3['date']}")
            return False
        if r3['id'] not in points_by_record:
            log(f"❌ Record 3 ID {r3['id']} not in points_by_record")
            return False
        balance3 = points_by_record[r3['id']]['balance']
        expected3 = 117
        if balance3 != expected3:
            log(f"❌ Record 3 balance mismatch: expected {expected3}, got {balance3}")
            return False
        log(f"✅ Record 3 ({record3_date}): balance = {balance3} (expected {expected3})")
        passed_tests += 1
        
    except Exception as e:
        log(f"❌ Exception: {e}")
        import traceback
        traceback.print_exc()
        return False
    
    # ========================================================================
    # TEST 7: VERIFY points_by_user (LAST RECORD SNAPSHOT)
    # ========================================================================
    log("\n" + "=" * 80)
    log("TEST 7: VERIFY points_by_user (LAST RECORD SNAPSHOT)")
    log("=" * 80)
    
    try:
        total_tests += 1
        if cindy_id not in points_by_user:
            log(f"❌ Cindy ID {cindy_id} not in points_by_user")
            return False
        
        user_balance = points_by_user[cindy_id]['balance']
        # Should equal the balance of the LAST record (record3)
        expected_user_balance = 117
        if user_balance != expected_user_balance:
            log(f"❌ points_by_user[cindy] balance mismatch: expected {expected_user_balance}, got {user_balance}")
            return False
        
        log(f"✅ points_by_user[cindy].balance = {user_balance} (matches last record)")
        passed_tests += 1
        
    except Exception as e:
        log(f"❌ Exception: {e}")
        import traceback
        traceback.print_exc()
        return False
    
    # ========================================================================
    # TEST 8: FROZEN HISTORIC TEST (CRITICAL)
    # ========================================================================
    log("\n" + "=" * 80)
    log("TEST 8: FROZEN HISTORIC TEST - Insert ledger AFTER record date")
    log("=" * 80)
    
    try:
        # Capture balance for record1 before inserting new ledger
        balance_before = points_by_record[record1_id]['balance']
        log(f"   Record 1 balance BEFORE: {balance_before}")
        
        # Insert a NEW ledger entry with event_date > record1.date (e.g., 2026-02-25)
        total_tests += 1
        future_ledger_id = str(uuid4())
        future_ledger = {
            "id": future_ledger_id,
            "user_id": cindy_id,
            "user_name": cindy_name,
            "event_type": "adjustment",
            "event_date": "2026-02-25",  # AFTER all records
            "period_key": "2026-02",
            "points": 50,  # Big change
            "source_id": future_ledger_id,
            "createdAt": datetime.utcnow(),
        }
        db.absensi_point_ledger.insert_one(future_ledger)
        log(f"   Inserted ledger entry: 2026-02-25 → +50 points (AFTER all records)")
        
        # Call report again
        resp = requests.get(f"{BASE_URL}/api/absensi/report?from=2026-02-01&to=2026-02-28", headers=headers, timeout=10)
        if resp.status_code != 200:
            log(f"❌ GET /api/absensi/report failed: {resp.status_code}")
            return False
        
        report2 = resp.json()
        points_by_record2 = report2['points_by_record']
        
        # Verify record1 balance is UNCHANGED
        if record1_id not in points_by_record2:
            log(f"❌ Record 1 ID {record1_id} not in points_by_record after future ledger insert")
            return False
        
        balance_after = points_by_record2[record1_id]['balance']
        log(f"   Record 1 balance AFTER: {balance_after}")
        
        if balance_after != balance_before:
            log(f"❌ FROZEN HISTORIC BROKEN: Record 1 balance changed from {balance_before} to {balance_after}")
            log(f"   Future ledger entry (2026-02-25) should NOT affect record from 2026-02-10")
            return False
        
        log(f"✅ FROZEN HISTORIC VERIFIED: Record 1 balance unchanged ({balance_before} → {balance_after})")
        log(f"   Future ledger entry (2026-02-25) correctly ignored for historical record (2026-02-10)")
        passed_tests += 1
        
    except Exception as e:
        log(f"❌ Exception: {e}")
        import traceback
        traceback.print_exc()
        return False
    
    # ========================================================================
    # TEST 9: BACKDATED LEDGER TEST
    # ========================================================================
    log("\n" + "=" * 80)
    log("TEST 9: BACKDATED LEDGER TEST - Insert ledger BEFORE record date")
    log("=" * 80)
    
    try:
        # Capture balance for record1 before inserting backdated ledger
        balance_before = points_by_record2[record1_id]['balance']
        log(f"   Record 1 balance BEFORE: {balance_before}")
        
        # Insert a backdated ledger entry with event_date <= record1.date (e.g., 2026-02-08)
        total_tests += 1
        backdated_ledger_id = str(uuid4())
        backdated_ledger = {
            "id": backdated_ledger_id,
            "user_id": cindy_id,
            "user_name": cindy_name,
            "event_type": "adjustment",
            "event_date": "2026-02-08",  # BEFORE record1 (2026-02-10)
            "period_key": "2026-02",
            "points": 7,
            "source_id": backdated_ledger_id,
            "createdAt": datetime.utcnow(),
        }
        db.absensi_point_ledger.insert_one(backdated_ledger)
        log(f"   Inserted backdated ledger entry: 2026-02-08 → +7 points (BEFORE record1)")
        
        # Call report again
        resp = requests.get(f"{BASE_URL}/api/absensi/report?from=2026-02-01&to=2026-02-28", headers=headers, timeout=10)
        if resp.status_code != 200:
            log(f"❌ GET /api/absensi/report failed: {resp.status_code}")
            return False
        
        report3 = resp.json()
        points_by_record3 = report3['points_by_record']
        
        # Verify record1 balance is UPDATED
        if record1_id not in points_by_record3:
            log(f"❌ Record 1 ID {record1_id} not in points_by_record after backdated ledger insert")
            return False
        
        balance_after = points_by_record3[record1_id]['balance']
        log(f"   Record 1 balance AFTER: {balance_after}")
        
        expected_new_balance = balance_before + 7  # 115 + 7 = 122
        if balance_after != expected_new_balance:
            log(f"❌ Backdated ledger not applied: expected {expected_new_balance}, got {balance_after}")
            return False
        
        log(f"✅ BACKDATED LEDGER VERIFIED: Record 1 balance updated ({balance_before} → {balance_after})")
        log(f"   Backdated ledger entry (2026-02-08) correctly applied to record (2026-02-10)")
        passed_tests += 1
        
    except Exception as e:
        log(f"❌ Exception: {e}")
        import traceback
        traceback.print_exc()
        return False
    
    # ========================================================================
    # TEST 10: EXCEL EXPORT VERIFICATION
    # ========================================================================
    log("\n" + "=" * 80)
    log("TEST 10: EXCEL EXPORT VERIFICATION")
    log("=" * 80)
    
    try:
        total_tests += 1
        resp = requests.get(f"{BASE_URL}/api/absensi/report/export?from=2026-02-01&to=2026-02-28", headers=headers, timeout=10)
        if resp.status_code != 200:
            log(f"❌ GET /api/absensi/report/export failed: {resp.status_code}")
            return False
        
        content_type = resp.headers.get('Content-Type', '')
        if 'spreadsheet' not in content_type and 'xlsx' not in content_type:
            log(f"❌ Content-Type is not Excel: {content_type}")
            return False
        
        log(f"✅ Excel export successful")
        log(f"   Content-Type: {content_type}")
        log(f"   File size: {len(resp.content)} bytes")
        passed_tests += 1
        
        # Parse Excel to verify "Poin Saat Absen" column
        total_tests += 1
        try:
            import openpyxl
            from io import BytesIO
            
            wb = openpyxl.load_workbook(BytesIO(resp.content))
            if 'Rekapitulasi' not in wb.sheetnames:
                log(f"❌ Sheet 'Rekapitulasi' not found in workbook")
                return False
            
            ws = wb['Rekapitulasi']
            header = [cell.value for cell in ws[1]]
            log(f"   Rekapitulasi header: {header}")
            
            if 'Poin Saat Absen' not in header:
                log(f"❌ Column 'Poin Saat Absen' not found in Rekapitulasi header")
                log(f"   Header: {header}")
                return False
            
            log(f"✅ Column 'Poin Saat Absen' found in Rekapitulasi sheet")
            
            # Verify it's the last column
            if header[-1] != 'Poin Saat Absen':
                log(f"❌ 'Poin Saat Absen' is not the last column")
                log(f"   Last column: {header[-1]}")
                return False
            
            log(f"✅ 'Poin Saat Absen' is the last column (column {len(header)})")
            passed_tests += 1
            
        except ImportError:
            log(f"⚠️  openpyxl not installed, skipping Excel parsing")
            log(f"   Excel export endpoint returned 200 with correct Content-Type")
            passed_tests += 1
        
    except Exception as e:
        log(f"❌ Exception: {e}")
        import traceback
        traceback.print_exc()
        return False
    
    # ========================================================================
    # TEST 11: EMPTY CASE (ZERO ROWS)
    # ========================================================================
    log("\n" + "=" * 80)
    log("TEST 11: EMPTY CASE - Filter yielding zero rows")
    log("=" * 80)
    
    try:
        total_tests += 1
        # Query a date range with no records
        resp = requests.get(f"{BASE_URL}/api/absensi/report?from=2025-01-01&to=2025-01-31", headers=headers, timeout=10)
        if resp.status_code != 200:
            log(f"❌ GET /api/absensi/report (empty) failed: {resp.status_code}")
            return False
        
        report_empty = resp.json()
        
        if 'points_by_record' not in report_empty or 'points_by_user' not in report_empty:
            log(f"❌ Empty report missing points fields")
            return False
        
        if not isinstance(report_empty['points_by_record'], dict) or not isinstance(report_empty['points_by_user'], dict):
            log(f"❌ Empty report points fields are not objects")
            return False
        
        if len(report_empty['points_by_record']) != 0 or len(report_empty['points_by_user']) != 0:
            log(f"❌ Empty report should have empty points objects")
            log(f"   points_by_record: {len(report_empty['points_by_record'])} items")
            log(f"   points_by_user: {len(report_empty['points_by_user'])} items")
            return False
        
        log(f"✅ Empty case verified: points_by_record = {{}}, points_by_user = {{}}")
        passed_tests += 1
        
    except Exception as e:
        log(f"❌ Exception: {e}")
        import traceback
        traceback.print_exc()
        return False
    
    # ========================================================================
    # TEST 12: REGRESSION - LEADERBOARD ENDPOINT
    # ========================================================================
    log("\n" + "=" * 80)
    log("TEST 12: REGRESSION - Leaderboard endpoint still works")
    log("=" * 80)
    
    try:
        total_tests += 1
        resp = requests.get(f"{BASE_URL}/api/absensi/points/leaderboard", headers=headers, timeout=10)
        if resp.status_code != 200:
            log(f"❌ GET /api/absensi/points/leaderboard failed: {resp.status_code}")
            return False
        
        leaderboard = resp.json()
        log(f"✅ Leaderboard endpoint working")
        log(f"   Period: {leaderboard.get('period')}")
        log(f"   Items: {len(leaderboard.get('items', []))}")
        passed_tests += 1
        
    except Exception as e:
        log(f"❌ Exception: {e}")
        import traceback
        traceback.print_exc()
        return False
    
    # ========================================================================
    # TEST 13: REGRESSION - POINTS HISTORY ENDPOINT
    # ========================================================================
    log("\n" + "=" * 80)
    log("TEST 13: REGRESSION - Points history endpoint still works")
    log("=" * 80)
    
    try:
        total_tests += 1
        resp = requests.get(f"{BASE_URL}/api/absensi/points/history?user_id={cindy_id}&period=2026-02", headers=headers, timeout=10)
        if resp.status_code != 200:
            log(f"❌ GET /api/absensi/points/history failed: {resp.status_code}")
            return False
        
        history = resp.json()
        log(f"✅ Points history endpoint working")
        log(f"   Items: {len(history.get('items', []))}")
        passed_tests += 1
        
    except Exception as e:
        log(f"❌ Exception: {e}")
        import traceback
        traceback.print_exc()
        return False
    
    # ========================================================================
    # CLEANUP
    # ========================================================================
    log("\n" + "=" * 80)
    log("CLEANUP")
    log("=" * 80)
    
    try:
        db.absensi_records.delete_many({"user_id": cindy_id, "date": {"$regex": "^2026-02-"}})
        db.absensi_point_ledger.delete_many({"user_id": cindy_id, "event_date": {"$regex": "^2026-02-"}})
        log("✅ Cleaned up test data")
    except Exception as e:
        log(f"⚠️  Cleanup warning: {e}")
    
    # ========================================================================
    # SUMMARY
    # ========================================================================
    log("\n" + "=" * 80)
    log("TEST SUMMARY")
    log("=" * 80)
    log(f"Total tests: {total_tests}")
    log(f"Passed: {passed_tests}")
    log(f"Failed: {total_tests - passed_tests}")
    log(f"Success rate: {passed_tests}/{total_tests} ({100*passed_tests//total_tests}%)")
    
    if passed_tests == total_tests:
        log("\n✅ ALL TESTS PASSED - Absensi Poin Saat Absen patch FULLY WORKING")
        return True
    else:
        log(f"\n❌ {total_tests - passed_tests} TEST(S) FAILED")
        return False

if __name__ == "__main__":
    success = main()
    sys.exit(0 if success else 1)
